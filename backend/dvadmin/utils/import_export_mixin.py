# -*- coding: utf-8 -*-
import datetime
from urllib.parse import quote

from django.db import transaction
from django.http import HttpResponse
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.table import Table, TableStyleInfo
from rest_framework.decorators import action
from rest_framework.request import Request

from dvadmin.utils.import_export import import_to_data
from dvadmin.utils.json_response import DetailResponse, SuccessResponse
from dvadmin.utils.request_util import get_verbose_name
from dvadmin.system.tasks import async_export_data
from dvadmin.system.models import DownloadCenter


class ImportSerializerMixin:
    """
    自定义导入模板、导入功能
    """

    # 导入字段
    import_field_dict = {}
    # 导入序列化器
    import_serializer_class = None
    # 表格表头最大宽度，默认50个字符
    export_column_width = 50

    def is_number(self,num):
        try:
            float(num)
            return True
        except (ValueError, TypeError):
            pass

        try:
            import unicodedata
            unicodedata.numeric(num)
            return True
        except (TypeError, ValueError):
            pass
        return False

    def get_string_len(self, string):
        """
        获取字符串最大长度
        :param string:
        :return:
        """
        length = 4
        if string is None:
            return length
        if self.is_number(string):
            return length
        for char in string:
            length += 2.1 if ord(char) > 256 else 1
        return round(length, 1) if length <= self.export_column_width else self.export_column_width

    @action(methods=['get','post'],detail=False)
    @transaction.atomic  # Django 事务,防止出错
    def import_data(self, request: Request, *args, **kwargs):
        """
        导入模板
        :param request:
        :param args:
        :param kwargs:
        :return:
        """
        assert self.import_field_dict, "'%s' " % self.__class__.__name__ + _("Please configure the corresponding export template fields.")
        # 导出模板
        if request.method == "GET":
            # 示例数据
            queryset = self.filter_queryset(self.get_queryset())
            # 导出excel 表
            _import_label = _("Import")
            _template_label = _("Template")
            response = HttpResponse(content_type="application/msexcel")
            response["Access-Control-Expose-Headers"] = f"Content-Disposition"
            response[
                "Content-Disposition"
            ] = f'attachment;filename={quote(str(f"{_import_label}{get_verbose_name(queryset)}{_template_label}.xlsx"))}'
            wb = Workbook()
            ws1 = wb.create_sheet("data", 1)
            ws1.sheet_state = "hidden"
            ws = wb.active
            row = get_column_letter(len(self.import_field_dict) + 1)
            column = 10
            header_data = [
                _("No."),
            ]
            validation_data_dict = {}
            # ★ 必填项从「导入序列化器」的字段 required 属性自动推断，
            #   不再靠人工把 * 写进标题 —— 人工写的星号经常与真实校验规则不一致
            #   （例如机房/环境/业务线模型上是 null=True，标了 * 会让人误以为必填）。
            required_keys = set()
            if self.import_serializer_class:
                try:
                    _ser = self.import_serializer_class()
                    required_keys = {n for n, f in _ser.fields.items()
                                     if getattr(f, "required", False)}
                except Exception:
                    required_keys = set()
            notes = []  # (列名, 是否必填, 可选值/填写要求)
            for index, (key, ele) in enumerate(self.import_field_dict.items()):
                options = []
                if isinstance(ele, dict):
                    # ★ 先剥掉标题里可能已经手写的 *，避免出现 "主机名**"
                    title = str(ele.get("title") or "").strip().rstrip("*").strip()
                    choices = ele.get("choices", {})
                    if choices.get("data"):
                        # data 形如 {'online': '在线'} ⇒ 下拉值用 key，说明里写成 "online=在线"
                        options = [f"{k}={v}" for k, v in choices.get("data").items()]
                        validation_data_dict[title] = list(choices.get("data").keys())
                    elif choices.get("queryset") and choices.get("values_name"):
                        options = list(
                            choices.get("queryset").values_list(choices.get("values_name"), flat=True)
                        )
                        validation_data_dict[title] = options
                    else:
                        header_data.append(title + ("*" if key in required_keys else ""))
                        notes.append((title, key in required_keys, ""))
                        continue
                    column_letter = get_column_letter(len(validation_data_dict))
                    dv = DataValidation(
                        type="list",
                        formula1=f"{quote_sheetname('data')}!${column_letter}$2:${column_letter}${len(validation_data_dict[title]) + 1}",
                        allow_blank=True,
                    )
                    ws.add_data_validation(dv)
                    dv.add(f"{get_column_letter(index + 2)}2:{get_column_letter(index + 2)}1048576")
                else:
                    title = str(ele).strip().rstrip("*").strip()
                    options = []
                header_data.append(title + ("*" if key in required_keys else ""))
                notes.append((title, key in required_keys, "、".join(str(o) for o in options)))
            # 添加数据列
            ws1.append(list(validation_data_dict.keys()))
            for index, validation_data in enumerate(validation_data_dict.values()):
                for inx, ele in enumerate(validation_data):
                    ws1[f"{get_column_letter(index + 1)}{inx + 2}"] = ele
            # 插入导出模板正式数据
            df_len_max = [self.get_string_len(ele) for ele in header_data]
            ws.append([str(x) if x is not None else "" for x in header_data])
            # 　更新列宽
            for index, width in enumerate(df_len_max):
                ws.column_dimensions[get_column_letter(index + 1)].width = width
            tab = Table(displayName="Table1", ref=f"A1:{row}{column}")  # 名称管理器
            style = TableStyleInfo(
                name="TableStyleLight11",
                showFirstColumn=True,
                showLastColumn=True,
                showRowStripes=True,
                showColumnStripes=True,
            )
            tab.tableStyleInfo = style
            ws.add_table(tab)
            # ★ 追加一页「填写说明」：
            #   ① 导入解析只读第一个 sheet（import_to_data 里 workbook[sheetnames[0]]），
            #      所以加这一页**不会**影响导入；
            #   ② 它解决的是「打开模板不知道哪些必填、枚举该填什么」这个问题。
            notes_ws = wb.create_sheet(_("填写说明"), 2)
            notes_ws.append([_("列名"), _("是否必填"), _("填写要求 / 可选值")])
            for title, is_required, options in notes:
                requirement = "必填，不能为空" if is_required else "可留空"
                if options:
                    requirement += "；可选值：" + options
                notes_ws.append([title + ("*" if is_required else ""),
                                 "是" if is_required else "否",
                                 requirement])
            notes_ws.append([])
            notes_ws.append([_("说明"), "", "带 * 的列为必填；关联列（机房/环境/业务线）"
                                               "必须填写「已存在」的名称，请先到对应管理页创建。"])
            for _c, _w in zip("ABC", (22, 10, 80)):
                notes_ws.column_dimensions[_c].width = _w
            wb.save(response)
            return response
        else:
            # 从excel中组织对应的数据结构，然后使用序列化器保存
            queryset = self.filter_queryset(self.get_queryset())
            # 获取多对多字段
            m2m_fields = [
                ele.name
                for ele in queryset.model._meta.get_fields()
                if hasattr(ele, "many_to_many") and ele.many_to_many == True
            ]
            import_field_dict = {'id': _("Update primary key (do not modify)"), **self.import_field_dict}
            data = import_to_data(request.data.get("url"), import_field_dict, m2m_fields)
            for ele in data:
                filter_dic = {'id':ele.get('id')}
                instance = filter_dic and queryset.filter(**filter_dic).first()
                # print(156,ele)
                serializer = self.import_serializer_class(instance, data=ele, request=request)
                serializer.is_valid(raise_exception=True)
                serializer.save()
            return DetailResponse(msg=_("Import successful"))

    @action(methods=['get'],detail=False)
    def update_template(self,request):
        queryset = self.filter_queryset(self.get_queryset())
        assert self.import_field_dict, "'%s' " % self.__class__.__name__ + _("Please configure the corresponding import template fields.")
        assert self.import_serializer_class, "'%s' " % self.__class__.__name__ + _("Please configure the corresponding import serializer.")
        data = self.import_serializer_class(queryset, many=True, request=request).data
        # 导出excel 表
        _export_label = _("Export")
        response = HttpResponse(content_type="application/msexcel")
        response["Access-Control-Expose-Headers"] = f"Content-Disposition"
        response["content-disposition"] = f'attachment;filename={quote(str(f"{_export_label}{get_verbose_name(queryset)}.xlsx"))}'
        wb = Workbook()
        ws1 = wb.create_sheet("data", 1)
        ws1.sheet_state = "hidden"
        ws = wb.active
        import_field_dict = {}
        header_data = [_("No."), _("Update primary key (do not modify)")]
        hidden_header = ["#","id"]
        #----设置选项----
        validation_data_dict = {}
        for index, item in enumerate(self.import_field_dict.items()):
            items = list(item)
            key = items[0]
            value = items[1]
            if isinstance(value, dict):
                header_data.append(value.get("title"))
                hidden_header.append(value.get('display'))
                choices = value.get("choices", {})
                if choices.get("data"):
                    data_list = []
                    data_list.extend(choices.get("data").keys())
                    validation_data_dict[value.get("title")] = data_list
                elif choices.get("queryset") and choices.get("values_name"):
                    data_list = choices.get("queryset").values_list(choices.get("values_name"), flat=True)
                    validation_data_dict[value.get("title")] = list(data_list)
                else:
                    continue
                column_letter = get_column_letter(len(validation_data_dict))
                dv = DataValidation(
                    type="list",
                    formula1=f"{quote_sheetname('data')}!${column_letter}$2:${column_letter}${len(validation_data_dict[value.get('title')]) + 1}",
                    allow_blank=True,
                )
                ws.add_data_validation(dv)
                dv.add(f"{get_column_letter(index + 3)}2:{get_column_letter(index + 3)}1048576")
            else:
                header_data.append(value)
                hidden_header.append(key)
        # 添加数据列
        ws1.append(list(validation_data_dict.keys()))
        for index, validation_data in enumerate(validation_data_dict.values()):
            for inx, ele in enumerate(validation_data):
                ws1[f"{get_column_letter(index + 1)}{inx + 2}"] = ele
        #--------
        df_len_max = [self.get_string_len(ele) for ele in header_data]
        row = get_column_letter(len(hidden_header) + 1)
        column = 1
        ws.append([str(x) if x is not None else "" for x in header_data])
        for index, results in enumerate(data):
            results_list = []
            for h_index, h_item in enumerate(hidden_header):
                for key, val in results.items():
                    if key == h_item:
                        if val is None or val == "":
                            results_list.append("")
                        elif isinstance(val,list):
                            results_list.append(str(val))
                        else:
                            results_list.append(val)
                        # 计算最大列宽度
                        if isinstance(val,str):
                            result_column_width = self.get_string_len(val)
                            if h_index != 0 and result_column_width > df_len_max[h_index]:
                                df_len_max[h_index] = result_column_width
            ws.append([index+1,*results_list])
            column += 1
        # 　更新列宽
        for index, width in enumerate(df_len_max):
            ws.column_dimensions[get_column_letter(index + 1)].width = width
        tab = Table(displayName="Table", ref=f"A1:{row}{column}")  # 名称管理器
        style = TableStyleInfo(
            name="TableStyleLight11",
            showFirstColumn=True,
            showLastColumn=True,
            showRowStripes=True,
            showColumnStripes=True,
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)
        wb.save(response)
        return response


class ExportSerializerMixin:
    """
    自定义导出功能
    """

    # 导出字段
    export_field_label = []
    # 导出序列化器
    export_serializer_class = None
    # 表格表头最大宽度，默认50个字符
    export_column_width = 50

    def is_number(self,num):
        try:
            float(num)
            return True
        except (ValueError, TypeError):
            pass

        try:
            import unicodedata
            unicodedata.numeric(num)
            return True
        except (TypeError, ValueError):
            pass
        return False

    def get_string_len(self, string):
        """
        获取字符串最大长度
        :param string:
        :return:
        """
        length = 4
        if string is None:
            return length
        if self.is_number(string):
            return length
        for char in string:
            length += 2.1 if ord(char) > 256 else 1
        return round(length, 1) if length <= self.export_column_width else self.export_column_width

    @action(methods=['get'],detail=False)
    def export_data(self, request: Request, *args, **kwargs):
        """
        导出功能
        :param request:
        :param args:
        :param kwargs:
        :return:
        """
        queryset = self.filter_queryset(self.get_queryset())
        assert self.export_field_label, "'%s' " % self.__class__.__name__ + _("Please configure the corresponding export template fields.")
        assert self.export_serializer_class, "'%s' " % self.__class__.__name__ + _("Please configure the corresponding export serializer.")
        data = self.export_serializer_class(queryset, many=True, request=request).data
        try:
            _export_label = _("Export")
            _data_export_label = _("Data export task")
            async_export_data.delay(
                data,
                str(f"{_export_label}{get_verbose_name(queryset)}-{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"),
                DownloadCenter.objects.create(creator=request.user, task_name=f'{get_verbose_name(queryset)}{_data_export_label}', dept_belong_id=request.user.dept_id).pk,
                self.export_field_label
            )
            return SuccessResponse(msg=_("Import task has been created. Please go to ‘Download Center’ to wait for download"))
        except:
            pass
        # 导出excel 表
        _export_label = _("Export")
        response = HttpResponse(content_type="application/msexcel")
        response["Access-Control-Expose-Headers"] = f"Content-Disposition"
        response["content-disposition"] = f'attachment;filename={quote(str(f"{_export_label}{get_verbose_name(queryset)}.xlsx"))}'
        wb = Workbook()
        ws = wb.active
        header_data = [_("No."), *self.export_field_label.values()]
        hidden_header = ["#", *self.export_field_label.keys()]
        df_len_max = [self.get_string_len(ele) for ele in header_data]
        row = get_column_letter(len(self.export_field_label) + 1)
        column = 1
        ws.append([str(x) if x is not None else "" for x in header_data])
        for index, results in enumerate(data):
            results_list = []
            for h_index, h_item in enumerate(hidden_header):
                for key,val in results.items():
                    if key == h_item:
                        if val is None or val=="":
                            results_list.append("")
                        else:
                            results_list.append(val)
                        # 计算最大列宽度
                        result_column_width = self.get_string_len(val)
                        if h_index !=0 and result_column_width > df_len_max[h_index]:
                            df_len_max[h_index] = result_column_width
            ws.append([index + 1, *results_list])
            column += 1
        # 　更新列宽
        for index, width in enumerate(df_len_max):
            ws.column_dimensions[get_column_letter(index + 1)].width = width
        tab = Table(displayName="Table", ref=f"A1:{row}{column}")  # 名称管理器
        style = TableStyleInfo(
            name="TableStyleLight11",
            showFirstColumn=True,
            showLastColumn=True,
            showRowStripes=True,
            showColumnStripes=True,
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)
        wb.save(response)
        return response
