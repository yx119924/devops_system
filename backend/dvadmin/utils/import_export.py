# -*- coding: utf-8 -*-
import logging
import os
import re
from datetime import datetime

import openpyxl
from django.conf import settings
from django.utils.translation import gettext_lazy as _

from dvadmin.utils.validator import CustomValidationError

logger = logging.getLogger(__name__)


def media_upload_candidates(file_url):
    """把「上传后回传的 file_url」解析成磁盘上**可能**的真实路径（按优先级返回列表）。

    ★★★ 为什么不能只用 settings.BASE_DIR（2026-09-30 内网实机事故）：
      上传落盘走 FileList.url（FileField，`MEDIA_ROOT = "media"`）——
      Django 的 `FileSystemStorage.location = os.path.abspath(settings.MEDIA_ROOT)`，
      而 MEDIA_ROOT 是**相对值** ⇒ 真实落盘位置 = `<进程 CWD>/media/files/…`。
      而这里原来写 `os.path.join(settings.BASE_DIR, file_url)` —— **两套口径**。
      只要 BASE_DIR ≠ CWD（或 BASE_DIR 本身被 conf/env.py 覆盖成相对值，内网正是这种），
      就必然 `[Errno 2] No such file or directory: 'backend/media/files/4/4/xxx.xlsx'`。
      ⇒ 权威口径必须与上传**完全同一套**：`os.path.abspath(MEDIA_ROOT)` + FileField 相对名。
        （`bastion/consumers.py`、`bastion/views/session.py` 也是直接用 MEDIA_ROOT，
          所以改用这个口径同时把它们也统一了。）
    """
    raw = str(file_url or "").strip().replace("\\", "/")
    if not raw:
        return []
    if os.path.isabs(raw):
        return [raw]
    # 归一化出 FileField 的相对名（去掉前端/DB 可能带上的各种前缀）
    bare = raw.lstrip("/")
    for prefix in ("backend/media/", "media/", "backend/"):
        if bare.startswith(prefix):
            bare = bare[len(prefix):]
            break
    media_root = os.path.abspath(str(getattr(settings, "MEDIA_ROOT", "media") or "media"))
    base_dir = str(getattr(settings, "BASE_DIR", "") or "")
    candidates = [
        os.path.join(media_root, bare),                    # ① 权威：与上传落盘完全一致
        os.path.join(media_root, raw.lstrip("/")),         # ② 少数环境把 media/ 也算进相对名
        (os.path.join(base_dir, raw) if base_dir else ""),  # ③ 旧逻辑（兼容 BASE_DIR 正确的环境）
        raw,                                               # ④ 直接相对 CWD
    ]
    result = []
    for c in candidates:
        if c and c not in result:
            result.append(c)
    return result


def _find_in_media_root(basename):
    """兜底：在 MEDIA_ROOT 内按文件名递归找（限深 4 层，只翻 media 目录）。

    ★ 只在「候选路径全落空」时才走这一步。存在的理由是：DB 里的 `file_url`
      与真实落盘布局可能因为历史原因对不上（前缀约定不同 / 换过 MEDIA_ROOT），
      这时按文件名兜一把比直接报错对用户有用得多。命中会打 WARNING 日志 ——
      **让偏差可见，但不要拦住用户**。
    """
    if not basename:
        return None
    root = os.path.abspath(str(getattr(settings, "MEDIA_ROOT", "media") or "media"))
    if not os.path.isdir(root):
        return None
    for dirpath, dirnames, filenames in os.walk(root):
        if os.path.relpath(dirpath, root).count(os.sep) >= 4:
            dirnames[:] = []
        if basename in filenames:
            return os.path.join(dirpath, basename)
    return None


def resolve_upload_path(file_url):
    """返回第一个**真实存在**的文件路径；都不存在则抛**可读**的 CustomValidationError。

    负面行为是刻意设计的：宁可给出一句能照着做的中文提示，也不要露出
    `[Errno 2] No such file or directory: <服务器绝对路径>`（既看不懂、又泄露目录结构）。
    """
    raw = str(file_url or "").strip().replace("\\", "/")
    basename = os.path.basename(raw)
    tried = media_upload_candidates(file_url)
    for path in tried:
        if os.path.isfile(path):
            return path
    fallback = _find_in_media_root(basename)
    if fallback:
        logger.warning(
            "导入文件在标准候选路径里没找到，已按文件名在 MEDIA_ROOT 内兜底命中：%s"
            "（file_url=%r，已尝试 %s）。建议排查 file_url 前缀约定与 MEDIA_ROOT 配置。",
            fallback, file_url, tried)
        return fallback
    logger.warning("导入文件定位失败：file_url=%r，已尝试 %s", file_url, tried)
    # 这里直接写中文：与 application/dispatch.py 中 XwOps 自定义报错的做法一致
    # （`_()` 走 locale 文件，新增 msgid 在 .po 里还没有 ⇒ 用户只会看到英文）。
    raise CustomValidationError(
        "导入文件「%s」在服务器上找不到（可能已被清理，或上传与导入的工作目录不一致）。"
        "请重新上传后再试；若反复出现，请在目标机执行 docker logs dvadmin3-django 查看已尝试的路径。"
        % basename
    )


def import_to_data(file_url, field_data, m2m_fields=None):
    """
    读取导入的excel文件
    :param file_url:
    :param field_data: 首行数据源
    :param m2m_fields: 多对多字段
    :return:
    """
    # 读取excel 文件（路径口径见 media_upload_candidates 的说明）
    file_path_dir = resolve_upload_path(file_url)
    workbook = openpyxl.load_workbook(file_path_dir)
    table = workbook[workbook.sheetnames[0]]
    theader = tuple(table.values)[0] #Excel的表头
    is_update = '更新主键(勿改)' in theader #是否导入更新
    if is_update is False: #不是更新时,删除id列
        field_data.pop('id')
    # 获取参数映射
    validation_data_dict = {}
    for key, value in field_data.items():
        if isinstance(value, dict):
            choices = value.get("choices", {})
            data_dict = {}
            if choices.get("data"):
                for k, v in choices.get("data").items():
                    data_dict[k] = v
            elif choices.get("queryset") and choices.get("values_name"):
                data_list = choices.get("queryset").values(choices.get("values_name"), "id")
                for ele in data_list:
                    data_dict[ele.get(choices.get("values_name"))] = ele.get("id")
            else:
                continue
            validation_data_dict[key] = data_dict
    # 创建一个空列表，存储Excel的数据
    tables = []
    for i, row in enumerate(range(table.max_row)):
        if i == 0:
            continue
        array = {}
        for index, item in enumerate(field_data.items()):
            items = list(item)
            key = items[0]
            values = items[1]
            value_type = 'str'
            if isinstance(values, dict):
                value_type = values.get('type','str')
            cell_value = table.cell(row=row + 1, column=index + 2).value
            if cell_value is None or cell_value=='':
                continue
            elif value_type == 'date':
                print(61, datetime.strptime(str(cell_value), '%Y-%m-%d %H:%M:%S').date())
                try:
                    cell_value = datetime.strptime(str(cell_value), '%Y-%m-%d %H:%M:%S').date()
                except:
                    raise CustomValidationError(_("Date format is incorrect"))
            elif value_type == 'datetime':
                cell_value = datetime.strptime(str(cell_value), '%Y-%m-%d %H:%M:%S')
            else:
            # 由于excel导入数字类型后，会出现数字加 .0 的，进行处理
                if type(cell_value) is float and str(cell_value).split(".")[1] == "0":
                    cell_value = int(str(cell_value).split(".")[0])
                elif type(cell_value) is str:
                    cell_value = cell_value.strip(" \t\n\r")
            if key in validation_data_dict:
                array[key] = validation_data_dict.get(key, {}).get(cell_value, None)
                if key in m2m_fields:
                    array[key] = list(
                        filter(
                            lambda x: x,
                            [
                                validation_data_dict.get(key, {}).get(value, None)
                                for value in re.split(r"[，；：|.,;:\s]\s*", cell_value)
                            ],
                        )
                    )
            else:
                array[key] = cell_value
        tables.append(array)
    data = [i for i in tables if len(i) != 0]
    return data
